import csv
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

source = Path(r"C:\Users\yunji\Documents\ChatGPT\AI\엑셀 파일\조선_AX_목업데이터 - 복사본.xlsx")
output = Path(r"D:\workspace\AXProject\outputs\조선_AX_목업데이터_CSV")
book = load_workbook(source, read_only=True, data_only=True)
files = sorted(output.glob("*.csv"))
assert len(files) == len(book.sheetnames) == 15

for path in files:
    name = path.stem.split("_", 1)[1]
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        actual = list(csv.reader(handle))
    expected = list(book[name].values)
    assert len(actual) == len(expected), (name, len(actual), len(expected))
    for row_number, (actual_row, expected_row) in enumerate(zip(actual, expected), 1):
        assert len(actual_row) == len(expected_row), (name, row_number)
        for column, (value, original) in enumerate(zip(actual_row, expected_row)):
            header = expected[0][column]
            if original is None:
                target = ""
            elif isinstance(original, datetime):
                target = original.strftime("%Y-%m-%d") if header == "plan_date" else original.strftime("%Y-%m-%dT%H:%M:00+09:00")
            elif isinstance(original, (int, float)) and row_number > 1:
                assert abs(float(value) - original) < 1e-9, (name, row_number, header, value, original)
                continue
            else:
                target = str(original)
            assert value == target, (name, row_number, header, value, target)
    print(f"{path.name}: {len(actual) - 1} rows verified")
