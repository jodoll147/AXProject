import datetime
import json
import pathlib

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / '엑셀 파일' / '02_조선_샘플데이터.xlsx'
SPEC = ROOT / '엑셀 파일' / '01_조선_데이터참고명세서.xlsx'
TEMPLATE = ROOT / 'work' / 'screen_template.html'
OUT = ROOT / 'outputs' / '조선_AX_이력탐색_화면.html'

TABLES = [
    'vessels', 'zones', 'processes', 'equipment', 'blocks', 'status_codes',
    'work_orders', 'production_plan', 'production_actual', 'equipment_intervals',
    'incidents', 'work_history', 'inspections', 'rework', 'equipment_snapshot',
]


def cell(v):
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%dT%H:%M')
    if isinstance(v, datetime.date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    data = {}
    for name in TABLES:
        rows = list(wb[name].iter_rows(values_only=True))
        head = rows[0]
        data[name] = [
            {k: cell(v) for k, v in zip(head, r) if v is not None}
            for r in rows[1:] if any(v is not None for v in r)
        ]
    spec_wb = openpyxl.load_workbook(SPEC, data_only=True)
    spec = {}
    for row in list(spec_wb['항목명세'].iter_rows(values_only=True))[2:]:
        table, field, label, dtype, need, key = row[:6]
        if table in TABLES and field:
            spec.setdefault(table, []).append({
                'field': field, 'label': label, 'type': dtype,
                'required': need == '필수', 'key': '' if key in (None, '—') else key,
            })
    data['_spec'] = spec
    payload = json.dumps(data, ensure_ascii=False, separators=(',', ':'))
    html = TEMPLATE.read_text(encoding='utf-8').replace('/*__DATA__*/null', payload)
    OUT.write_text(html, encoding='utf-8')
    print(len(html), {k: len(v) for k, v in data.items()})


if __name__ == '__main__':
    main()
