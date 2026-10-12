from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path(__file__).resolve().parents[1]/'outputs'
def shade(cell,color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); tcPr.append(shd)
def set_cell(cell,text,bold=False,color=None,size=9):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(size); r.font.name='Malgun Gothic'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    if color:r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def setup(doc,title,subtitle):
    sec=doc.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
    styles=doc.styles
    for nm in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        st=styles[nm];st.font.name='Malgun Gothic';st._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    styles['Normal'].font.size=Pt(10);styles['Normal'].font.color.rgb=RGBColor(40,51,66)
    styles['Title'].font.size=Pt(25);styles['Title'].font.bold=True;styles['Title'].font.color.rgb=RGBColor(19,46,72)
    styles['Heading 1'].font.color.rgb=RGBColor(22,86,112);styles['Heading 1'].font.size=Pt(16)
    styles['Heading 2'].font.color.rgb=RGBColor(35,106,128);styles['Heading 2'].font.size=Pt(12)
    p=doc.add_paragraph(style='Title');p.add_run(title)
    p=doc.add_paragraph();r=p.add_run(subtitle);r.italic=True;r.font.color.rgb=RGBColor(95,112,130)
    doc.add_paragraph('문서 상태: 첨부 요구 메일과 데이터 참고 명세, 샘플 데이터 기준 초안 | 기준일: 2026-10-07')
def para(doc,text):doc.add_paragraph(text)
def bullets(doc,items):
    for x in items:doc.add_paragraph(x,style='List Bullet')
def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Table Grid'
    for i,h in enumerate(headers):set_cell(t.rows[0].cells[i],h,True,'FFFFFF',9);shade(t.rows[0].cells[i],'195A73')
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,x in enumerate(row):set_cell(cells[i],x,size=8.5);shade(cells[i],'F1F6F8' if ri%2==0 else 'FFFFFF')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths):row.cells[i].width=Inches(w)
    doc.add_paragraph()

# Requirements specification
d=Document();setup(d,'조선 AX 이력 탐색 요구사항 명세서','단일 화면 기반 공정·설비·품질 이력 조회 시스템')
d.add_heading('1. 문서 목적과 범위',1);para(d,'본 문서는 조선 블록 생산 샘플 데이터를 활용해 기준시각별 공정 상태와 작업·설비·사건 이력을 확인하는 시스템의 요구사항을 정의한다. 목표는 실시간 관제가 아니라 과거 시점을 재구성하고 이력의 근거를 추적하는 것이다. 실제 MES/PLC 연계, 실제 공장 축척도, 운영 계정 정책은 제공되지 않아 본 범위에서는 설계 전제로 확정하지 않는다.')
d.add_heading('2. 이해관계자와 우선순위',1);table(d,['우선순위','요청 항목','요구사항 반영'],[
['1순위','데이터·양식, 구현·평가 우선순위','샘플 CSV/엑셀 적재와 필수값·키·단위·시간·공정순서 검증 후 기본 조회 및 데이터 연결을 우선한다.'],
['2순위','기존 제품 화면 예시, 업무 흐름, 이력·이상 시나리오','기존 제품 화면은 미제공. 선박-블록-작업지시-실적-사건-검사 관계를 단일 화면에서 탐색한다.'],
['3순위','공정·배치도, AI 운영요약, 구현 고려','6개 구역 개념도 제공. AI 요약은 계산 결과와 사건 ID를 인용하며, CSV 오류·역할·시간/수량 정책은 운영 합의 대상으로 둔다.']])
d.add_heading('3. 사용자와 주요 업무',1);bullets(d,['운영 담당자는 선박·블록·공정·기간을 선택해 기준시각의 작업 상태를 확인한다.','공정/설비 담당자는 사건 발생·조치·복구·재개·완료 순서를 확인한다.','품질 담당자는 검사 불합격, 재작업 완료, 재검사 결과를 구분해 확인한다.','관리자는 계획 대비 실적과 지연을 이후 단계에서 확인한다.'])
d.add_heading('4. 기능 요구사항',1);table(d,['ID','요구사항','우선순위','수용 기준'],[
['FR-01','CSV/엑셀 데이터 입력 및 스키마 검증','P1','필수 컬럼, 중복 기본키, 참조 무결성, 단위, 시각 형식, 공정 순서 오류를 행·컬럼 단위로 알린다.'],
['FR-02','기준시각 선택 및 이력 재생','P1','시작/종료 범위 안에서 기준시각을 이동하고 30분 단위 이전·다음 및 재생/일시정지를 제공한다.'],
['FR-03','상태 요약과 필터','P1','선박 선택과 완료/진행/차단/계획/미해결 사건 집계를 제공한다. 모든 값은 기준시각 이전 데이터로 계산한다.'],
['FR-04','공정·구역 배치 탐색','P2','Z01~Z06과 P01~P06/EQ01~EQ06을 표시하고 Z03에 EQ07/EQ08을 보조 자원으로 표시한다. 개념도임을 명시한다.'],
['FR-05','작업지시·블록 이력 상세','P2','블록, 작업지시, 공정, 설비, 기준시각 상태, 진척, 관련 이력/사건을 확인한다.'],
['FR-06','사건 분류와 생명주기','P2','설비 고장, 자재 대기, 품질 보류를 분리하고 발생/조치/복구/재개 이력을 시간순으로 보여준다.'],
['FR-07','검사·재작업·재검사 추적','P2','불합격 수량과 재작업/재검사를 연결하며 재검 합격 전 합격수량을 올리지 않는다.'],
['FR-08','계획 대비 실적·지연','P3','동일 공정과 수량 단위를 비교한다. 계획 시작 전 작업은 지연으로 계산하지 않는다.'],
['FR-09','근거 기반 운영 요약','확장','요약 문장에 계산 시각, 관련 지표, 사건/작업지시 ID를 포함하고 입력에 없는 원인을 단정하지 않는다.']])
d.add_heading('5. 시나리오 기반 수용 기준',1);table(d,['시나리오','기대 결과'],[
['B001 정상 흐름','WO001-01~06의 순차 흐름과 2026-09-12 12:00 이송 완료를 확인한다.'],
['B004 설비 고장 복구','INC001의 10:00 정지, 11:30 복구, 90분 중단, 13:30 완료를 별도 이벤트로 표시한다.'],
['B006 자재 대기 해소','INC002를 설비 고장으로 분류하지 않고 10:00 대기, 11:00 재개, 75%, 예상 종료 13:00을 표시한다.'],
['B008 미해결 고장','INC003의 10:30 EQ03 알람, 62.5%와 복구 미확인 상태를 유지한다.'],
['B009 미해결 자재 대기','INC004의 09:00 대기, 25%, 작업 차단/설비 유휴를 표시한다.'],
['B007 품질 재작업','검사 10개 중 2개 불합격, 10:30~11:30 재작업 후 재검사 대기 및 합격수량 8을 표시한다.']])
d.add_heading('6. 데이터·계산 규칙',1);bullets(d,['기준시각 이후의 이벤트와 실적은 집계에서 제외한다.','accepted_qty_delta는 작업지시 단위로 사전 집계해 조인 증폭을 막는다.','설비 스냅샷은 2026-09-16 12:00에만 사용하고 다른 시점은 [start_at, end_at) 구간으로 계산한다.','사건은 resolved_at이 기준시각 이하일 때만 해소 상태다.','미래 계획은 지연이 아니며 계획/실적의 공정과 수량 단위를 맞춘다.'])
d.add_heading('7. 비기능 요구사항과 범위 제한',1);bullets(d,['화면은 다크 모드, 단일 화면, 반응형 레이아웃을 제공한다.','같은 입력 데이터와 기준시각은 같은 결과를 내도록 계산 규칙을 명시한다.','수정 이력, 사용자 권한, 접근 통제, 장기 보존기간은 운영 정책 확정 후 상세화한다.','샘플은 교육/설계용이다. 실시간 PLC/MES 연결, 실제 공장 배치, 운영 SLA는 검증되지 않았다.'])
d.add_heading('8. 샘플 데이터 범위',1);para(d,'원천 데이터 15종: vessels, zones, processes, equipment, blocks, status_codes, work_orders, production_plan, production_actual, equipment_intervals, incidents, work_history, inspections, rework, equipment_snapshot. 확인된 표본 규모는 작업지시/계획 72, 생산실적 133, 설비 구간 201, 사건 5, 작업이력 94, 검사 7, 재작업 1행이다. 목업 요약 시트의 KPI 불일치와 #DIV/0!는 원천 사실로 취급하지 않는다.')
d.save(OUT/'조선_AX_요구사항명세서.docx')

# Software design specification
d=Document();setup(d,'조선 AX 이력 탐색 소프트웨어 설계서','요구사항을 데이터·화면·처리 구조로 구체화한 설계')
d.add_heading('1. 설계 개요',1);para(d,'제공 샘플을 기준으로 한 단일 페이지 이력 탐색기다. 사용자가 기준시각을 바꾸면 그 시각까지 발생한 실적·설비 구간·사건·검사 이력을 재계산한다. 현재값 자동 갱신을 전제로 하지 않는다. 구현 단계는 데이터 입력/검증과 시점 조회부터 시작한다.')
d.add_heading('2. 논리 구성',1);table(d,['구성','책임','설계 메모'],[
['입력 계층','CSV/엑셀 파일 수신','파일/시트/컬럼 확인, 원본 적재와 오류 리포트'],['정합성 계층','키·관계·단위·시각·순서 검사','행 단위 오류를 보존해 수정 가능하게 한다.'],['도메인 계층','선박-블록-작업지시-실적-사건/검사 연결','as_of 필터와 사건 생명주기 계산'],['조회 계층','필터·요약·타임라인·상세 조회','집계 후 조인, 상태 재구성'],['표현 계층','단일 다크 화면','시점 컨트롤, 6구역 개념도, 작업/사건 목록']])
d.add_heading('3. 데이터 모델 및 ERD',1);para(d,'ERD는 원천 15개 테이블을 중심으로 하며 계정/감사로그는 운영 확장 항목이다. 실제 FK와 컬럼 타입은 원천 명세 파일로 최종 검증한다.')
erd='''vessels ||--o{ blocks : contains
zones ||--o{ processes : hosts
processes ||--o{ equipment : uses
blocks ||--o{ work_orders : processed_by
processes ||--o{ work_orders : defines
equipment o|--o{ work_orders : assigned
work_orders ||--o{ production_plan : planned
work_orders ||--o{ production_actual : recorded
equipment ||--o{ equipment_intervals : state_history
work_orders ||--o{ incidents : affected
equipment o|--o{ incidents : related
work_orders ||--o{ work_history : timeline
work_orders ||--o{ inspections : inspected
inspections ||--o{ rework : may_create
status_codes ||--o{ work_orders : status
equipment ||--o{ equipment_snapshot : snapshot'''
for line in erd.splitlines():p=d.add_paragraph();r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(9);r.font.color.rgb=RGBColor(30,90,110)
table(d,['엔터티','키/주요 속성','관계'],[
['vessels','vessel_id PK, vessel_name','1:N blocks'],['zones','zone_id PK, zone_name','1:N processes'],['processes','process_id PK, zone_id FK, sequence_no','1:N equipment/work_orders'],['equipment','equipment_id PK, process_id FK, equipment_type','1:N 상태구간/사건'],['blocks','block_id PK, vessel_id FK, block_type','1:N work_orders'],['work_orders','work_order_id PK, block_id/process_id FK, sequence_no','1:N 계획·실적·사건·이력·검사'],['production_plan','plan_id PK, work_order_id FK, planned_start/end, qty/unit','작업지시 계획'],['production_actual','actual_id PK, work_order_id FK, occurred_at, accepted_qty_delta','작업지시 실적 이벤트'],['equipment_intervals','interval_id PK, equipment_id FK, start_at/end_at, state','설비 시간 구간'],['incidents','incident_id PK, work_order_id FK, type, occurred_at/resolved_at','작업/설비 사건'],['work_history','history_id PK, work_order_id FK, event_at, event_type','작업 타임라인'],['inspections/rework','inspection_id/rework_id PK, work_order_id/inspection_id FK','검사에서 재작업 및 재검사']])
d.add_heading('4. 시점 조회 처리',1);table(d,['단계','처리'],[
['1. 기준시각 고정','입력 as_of 및 선택 필터를 요청 컨텍스트로 확정한다.'],['2. 이벤트 필터','occurred_at/event_at/start_at <= as_of만 대상. 종료 구간은 end_at이 null이거나 as_of보다 뒤인 경우 유효.'],['3. 상태 계산','작업은 실적/이력, 설비는 유효한 구간, 사건은 resolved_at <= as_of 기준으로 상태 산출.'],['4. 수량 집계','production_actual을 work_order_id별 먼저 합산하고 검사/재작업 결과를 반영한다.'],['5. 화면 조회','KPI, 구역 상태, 작업 목록, 사건 타임라인을 동일한 as_of 스냅샷으로 렌더링한다.']])
d.add_heading('5. 화면 설계',1);para(d,'한 페이지에 기준시각 조작, KPI, 공정 구역, 사건/작업 이력을 배치한다. 기준시각 슬라이더와 30분 단위 이동/재생을 제공한다. 구역은 2행 3열의 개념도이며 실제 공장 좌표를 암시하지 않는다. 상세는 같은 화면의 모달/패널로 열어 맥락을 유지한다.')
table(d,['영역','내용','상호작용'],[['헤더','서비스명, 이력 조회 모드, 설명','현재 실시간으로 오인되지 않도록 이력 조회 표기'],['시간 컨트롤','기준시각, 이전/다음, 재생, 슬라이더','30분 간격 이동'],['요약 KPI','완료/진행/차단/계획/미해결 사건','기준시각에 따라 재계산'],['공정 개념도','Z01~Z06, P01~P06, 장비 상태','Z03에 EQ07/EQ08 병기'],['작업·사건','사건 타임라인 및 작업지시 표','행 선택 시 상세/연관 이벤트'],['우선순위 표시','데이터 검증 → 이력 조회 → 지연 분석 → 근거 요약','개발 순서를 전달']])
d.add_heading('6. 공정 배치와 상태 표현',1);table(d,['구역','공정/주 설비','표현 규칙'],[['Z01','P01 강재 절단 / EQ01','기준시각 작업/상태'],['Z02','P02 소조립·취부 / EQ02','동일'],['Z03','P03 블록 용접 / EQ03, EQ07, EQ08','EQ03 주 장비, EQ07 보수·EQ08 예비'],['Z04','P04 치수·용접 검사 / EQ04','검사/재검사 사건 연계'],['Z05','P05 표면처리·도장 / EQ05','자재 대기는 설비 이상과 다른 색/유형'],['Z06','P06 블록 이송 / EQ06','완료 이력 연결']])
d.add_heading('7. 예외 및 데이터 품질',1);bullets(d,['중복 기본키/없는 참조키/필수값 누락은 적재 오류로 격리한다.','시간 역전, 음수 수량, 비정상 공정 순서는 오류 또는 검토 경고로 구분한다.','사건 원인이 비어 있으면 미확인으로 표시한다. 원인 추정 문장은 생성하지 않는다.','목업 요약 시트는 상세 원천과 대조하며 KPI 차이가 있으면 discrepancy로 기록한다.','실제 PLC/MES 연계가 시작되면 외부 timestamp 기준과 중복 이벤트 idempotency 규칙을 추가한다.'])
d.add_heading('8. 보안·운영·확장',1);para(d,'현재는 로컬 샘플 화면과 설계 수준이다. 운영 적용 시 사용자/역할, 감사로그, 파일 보존·삭제, 개인정보 여부, 데이터 내보내기 권한, 암호화와 배포 환경을 별도 결정해야 한다. 요약 생성 기능은 규칙 기반 지표를 우선하고, 추후 AI를 붙일 때도 지표/사건 근거를 함께 기록한다.')
d.add_heading('9. 검증 기준과 한계',1);para(d,'첨부 시나리오 6종을 기준시각별로 재현한다. 실제 운영 데이터 연계, 공장 도면 정합성, 원천 상태코드 사전은 제공 범위 밖이다. 샘플 표본 수치는 운영 용량 산정이나 성능 목표로 사용하지 않는다.')
d.save(OUT/'조선_AX_소프트웨어설계서.docx')
print('created',*(str(p) for p in OUT.iterdir()))
